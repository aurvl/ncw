"""Create the BNC2 supplementary material from the current accounting protocol."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[2] / "Supplementary Material.docx"
BLACK = RGBColor(0, 0, 0)
GRAY = "D9D9D9"


def set_cell_shading(cell: object, fill: str) -> None:
    """Apply a solid cell background colour."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell: object, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    """Add readable internal padding to a table cell."""
    tc_mar = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        cell._tc.get_or_add_tcPr().append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    """Add a compact, readable table with a dark header row."""
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = label
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "404040")
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[index])
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(8)
    document.add_paragraph()


def add_paragraph(document: Document, text: str, style: str | None = None) -> None:
    """Add one body paragraph with the project-wide spacing rules."""
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    paragraph.add_run(text)


def configure_document(document: Document) -> None:
    """Apply a restrained research-paper layout."""
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    styles = document.styles
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = BLACK
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.size = Pt(16)
    styles["Heading 1"].font.size = Pt(13)
    styles["Heading 2"].font.size = Pt(11)
    styles["Heading 3"].font.size = Pt(10)


def build_document() -> Document:
    """Build the new supplementary material in the order used by the paper."""
    document = Document()
    configure_document(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Supplementary Material")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Integrating forest and ocean carbon sinks into the natural carbon wealth of nations")

    document.add_heading("S1 Accounting object", level=1)
    add_paragraph(
        document,
        "Green Carbon Wealth is the annual welfare contribution of carbon uptake by forests. "
        "It is not a forest carbon stock, a market price, a finance requirement, or an estimate of additional mitigation. "
        "For each accounting unit, annual forest uptake in tonnes of carbon dioxide is valued with the Global Social Cost of Carbon (GSCC). "
        "Natural Carbon Wealth adds the harmonised blue component after the mangrove overlap is removed.",
    )
    add_paragraph(
        document,
        "All values use constant 2015 US dollars. The ecological reference period is 2010-2019. "
        "Population, GDP, debt, and fossil and industrial CO2 emissions refer to 2023 or the nearest available year. "
        "The emissions series excludes land-use change. This separation avoids double counting the deforestation analysis.",
    )

    document.add_heading("S2 Data, accounting units, and harmonisation", level=1)
    add_table(
        document,
        ["Input", "Source and period", "Use in the analysis"],
        [
            ["Forest area and mangroves", "FAO FRA 2025", "Country forest area and mangrove overlap"],
            ["Biome shares", "GEZ 2010, EEZ land union v4, JRC Global Forest Cover 2020", "Forest area by tropical, temperate, and boreal Pan biome"],
            ["Forest uptake rates", "Pan et al. 2024, 2010-2019", "Biome-specific annual uptake rates"],
            ["Blue Carbon Wealth", "Published BCW draw set", "Coastal and ocean components, GSCC draws, socioeconomic fields"],
            ["Deforestation", "FAO FRA 2025", "Reported annual forest conversion over 2010-2019"],
            ["Socioeconomic indicators", "BCW harmonised dataset, 2023", "GDP, population, debt, and fossil and industrial CO2 emissions"],
        ],
    )
    add_paragraph(
        document,
        "The accounting units follow FRA countries and territories. BCW rows are mapped to those units with a documented crosswalk. "
        "BCW sub-territories are regrouped where they belong to one FRA unit. Population and GDP are never summed across parts. "
        "Taiwan is included in China, consistent with FRA reporting. BCW-only territories and joint-regime rows remain separate where no FRA country match exists.",
    )

    document.add_heading("S3 Forest uptake estimate", level=1)
    add_paragraph(
        document,
        "The main estimate uses Route B. Forest area from FRA is distributed across biomes with a forest mask. "
        "The tropical and subtropical shares are combined because Pan et al. include subtropical forests in their tropical biome. "
        "Mangrove area is removed from the forest component before the forest uptake calculation.",
    )
    add_table(
        document,
        ["Pan biome", "Mean annual uptake rate", "Standard deviation", "Treatment"],
        [
            ["Tropical, including subtropical", "1.32 tC ha-1 yr-1", "0.21", "One shared tropical rate"],
            ["Temperate", "0.87 tC ha-1 yr-1", "0.06", "One shared temperate rate"],
            ["Boreal", "0.28 tC ha-1 yr-1", "0.03", "One shared boreal rate"],
        ],
    )
    add_paragraph(
        document,
        "This approach does not map FAO naturally regenerated and planted forest classes to Pan intact and regrowth classes. "
        "Such a mapping would be unsupported. The analysis therefore uses one Pan-consistent rate for each of the three biomes. "
        "The pipeline checks that the derived rates reproduce the published Pan biome fluxes and reported area-based values.",
    )

    document.add_heading("S4 Mangroves and the forest-ocean overlap", level=1)
    add_paragraph(
        document,
        "Mangroves are counted in FAO forest area and in the blue-carbon framework. They are removed from GCW and retained in BCW. "
        "This rule avoids double counting in NCW. Units with FRA mangroves that are absent from the published BCW file receive a mangrove estimate from the BCW framework. "
        "The gap-fill uses FRA mangrove area and the BCW mangrove-rate draws. Indonesia, Cuba, Senegal, and several very small units require this treatment.",
    )

    document.add_heading("S5 Uncertainty and aggregation", level=1)
    add_paragraph(
        document,
        "The analysis uses 20,000 Monte Carlo draws. Each biome rate follows a normal distribution truncated at zero. "
        "One draw per biome is shared by all countries. GCW uses the same GSCC draw as BCW. NCW totals are summed draw by draw. "
        "Reported intervals are the 5th and 95th percentiles of the resulting distribution. Standard deviations are not added across components.",
    )
    add_paragraph(
        document,
        "Areas are treated as fixed. This choice means that the uncertainty intervals reflect uptake rates and the GSCC, rather than uncertainty in national forest area or biome assignment.",
    )

    document.add_heading("S6 Statistical analyses and supplementary results", level=1)
    add_paragraph(
        document,
        "R1 reports absolute GCW, GCW relative to GDP, concentration, and the correlation between GCW and GDP. "
        "The group comparison uses a Kruskal-Wallis test. Pairwise comparisons use Mann-Whitney tests with Holm adjustment. "
        "Per-capita indicators are reported in the supplementary country table because they do not distinguish the main green-carbon narrative across groups.",
    )
    add_paragraph(
        document,
        "R2 classifies units as green-dominant when GCW represents at least 66.7% of NCW and blue-dominant when it represents at most 33.3%. "
        "It compares the net balance from forests alone with the balance after the ocean component. The debt-carbon quadrant uses the same debt threshold as the BCW analysis. "
        "The threshold must be checked against the final BCW documentation before submission.",
    )
    add_paragraph(
        document,
        "Supplementary Table S1 provides country-level GCW, BCW, NCW, uncertainty intervals, and GDP-normalised values. "
        "Supplementary Table S2 lists per-capita results. Supplementary Figure S1 shows the debt-carbon quadrants. "
        "Supplementary Figure S2 gives the country rankings and rank stability across Monte Carlo draws.",
    )

    document.add_heading("S7 Erosion of forest sink capacity", level=1)
    add_paragraph(
        document,
        "R3 defines a loss as forest area converted each year multiplied by the country biome shares and the Pan uptake rates. "
        "The result is the annual sequestration capacity that is removed by deforestation. It is not an estimate of direct land-use emissions. "
        "The calculation retains missing reports as missing values. It does not treat them as zero loss.",
    )
    add_paragraph(
        document,
        "FRA reports deforestation for 139 of 236 accounting units. These units represent 79% of global GCW. "
        "The global loss totals therefore describe the reporting units only. The allocation of reported deforestation across national biomes is an assumption. "
        "It may understate tropical losses where conversion is concentrated in tropical forest within a multi-biome country.",
    )
    add_paragraph(
        document,
        "The trend analysis uses the five FRA reporting intervals. It reports Theil-Sen slopes and Mann-Kendall tests. "
        "Five time points provide limited power. Country trend results should therefore be read as descriptive evidence.",
    )

    document.add_heading("S8 Compensation scenarios", level=1)
    add_paragraph(
        document,
        "R4 distinguishes welfare value from replacement cost. The GSCC values the welfare contribution that is lost when a sink weakens. "
        "An abatement price estimates the cost of replacing the lost tonnes if the climate objective is held constant. "
        "The comparison does not model a carbon market or a marginal abatement-cost curve.",
    )
    add_paragraph(
        document,
        "The current calculations use illustrative prices of 50, 150, and 400 constant 2015 US dollars per tonne of carbon dioxide. "
        "These values support a sensitivity analysis only. The final manuscript must identify the chosen AR6 scenario, price year, and conversion method. "
        "The allocation analysis compares domestic compensation with a shared effort proportional to GDP.",
    )
    add_paragraph(
        document,
        "The ocean component has no observed national loss series in this analysis. The supplementary scenarios apply 5% and 10% reductions to BCW. "
        "They illustrate the scale of an ocean sink decline. They do not estimate observed ocean degradation.",
    )

    document.add_heading("S9 Limitations", level=1)
    add_paragraph(
        document,
        "The study applies biome-average uptake rates to national forest areas. It does not capture national variation in forest age, disturbance, management, or climate response. "
        "The FAO and Pan forest definitions also differ. The analysis documents this mismatch and does not infer a direct correspondence between their classes.",
    )
    add_paragraph(
        document,
        "The ecological and socioeconomic series use different reference periods. Forest uptake refers to 2010-2019, whereas socioeconomic indicators refer to 2023. "
        "This temporal mismatch supports a comparative accounting exercise but limits causal interpretation. The deforestation estimates also have incomplete reporting coverage.",
    )
    add_paragraph(
        document,
        "GCW and NCW are annual welfare contributions. They are not estimates of carbon stocks, ownership, payment obligations, or additional mitigation potential. "
        "The compensation results depend on the selected abatement price and allocation rule. They are scenario results, not observed transfers.",
    )

    document.add_heading("S10 Data and code availability", level=1)
    add_paragraph(
        document,
        "The reproducible pipeline, input registry, validation checks, country-level outputs, and manuscript artefact builder are stored in the project repository. "
        "The notebook `04_analyses_R1_R4.ipynb` contains the analytical workflow. The notebook `05_manuscript_tables_figures.ipynb` regenerates the main-text tables and figures without overwriting the working analysis outputs.",
    )
    return document


def main() -> None:
    """Write the supplementary material to the project root."""
    document = build_document()
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
