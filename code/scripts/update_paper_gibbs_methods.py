"""Replace the obsolete Pan-based forest-rate methods in NEXT/PAPER.docx.

The script backs up the source, changes only method paragraphs and adds the
Gibbs reference. Results and introduction require a separate numeric revision.
"""

from __future__ import annotations

from pathlib import Path
import hashlib
import shutil
import zipfile

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parents[2]
PAPER = ROOT.parent / "NEXT" / "PAPER.docx"
BACKUP = ROOT / "results" / "validation" / "PAPER_before_Gibbs_methods.docx"


REPLACEMENTS = {
    134: (
        "We estimate an annual flow contribution to inclusive wealth. Green Carbon Wealth (GCW) is the annual welfare value assigned to carbon added to above- and below-ground living forest biomass, using a global welfare-based shadow price of carbon. This is the gross-removal component of the Gibbs/GFW forest carbon model, not its net forest ecosystem balance. Blue Carbon Wealth (BCW) is the corresponding published account for coastal blue carbon ecosystems and the oceanic biological carbon pump (Hilmi et al., 2026). Natural Carbon Wealth (NCW) combines them under the stated mangrove-area convention."
    ),
    141: (
        "In this study, forest carbon sequestration refers operationally to the modeled annual incorporation of atmospheric carbon into living above- and below-ground biomass. Gibbs et al. (2025) estimate this gross-removal component separately from emissions caused by forest disturbances. Carbon subsequently leaves different living tissues over different timescales; our rate does not measure residence time or the net carbon balance of the entire forest ecosystem."
    ),
    145: (
        "Forest areas come from FRA 2025, which applies the FAO definition of forest: land spanning more than 0.5 hectares with trees higher than 5 meters and canopy cover above 10 per cent, excluding land predominantly under agricultural or urban use. Forest carbon also occurs in dead wood, litter and soils, but these reservoirs are outside the forest-flow indicator valued here. GCW applies mean annual gross additions to living above- and below-ground biomass to the allocated forest area. Reported areas are official national submissions for 194 of 236 units and FAO desk-study estimates for the remaining 42."
    ),
    147: (
        "The current national calculation deducts the FRA-reported mangrove category from the allocated tropical forest area before applying the forest rates, while BCW retains mangrove sediment burial. This area adjustment is a proxy: FRA's separate mangrove category includes both forest and other wooded land, and the part already included in national forest area is not identified for every unit. We flag rather than conceal cases where the reported mangrove area exceeds the allocated tropical forest area. The independent rate derivation also excludes pixels classified as mangroves by the Gibbs/GFW source model."
    ),
    148: "Gross living-biomass addition rates from Gibbs/GFW",
    149: (
        "Gibbs et al. (2025, Table 4) report average annual gross-removal totals for 2001-2023 in the climate domains of their model. Expressed as positive magnitudes, these are 7.1 ± 7.6 Gt CO₂ yr⁻¹ in tropical forests, 1.7 ± 0.56 in subtropical forests, 3.1 ± 0.55 in temperate forests and 2.5 ± 0.95 in boreal forests; the published global total is 14.5 ± 7.7 Gt CO₂ yr⁻¹. The ± values are the authors' standard deviations. Their table prints removals with negative signs under an emissions-positive convention. These totals are not per-hectare rates, and the model's climate-domain boundaries do not necessarily coincide with our FAO GEZ domains."
    ),
    150: (
        "For each valid source pixel p, rₚ = Rₚ × (12/44) / 23, where Rₚ is cumulative gross removal in Mg CO₂ ha⁻¹ over 2001-2023 and rₚ is mean annual addition in tC ha⁻¹ yr⁻¹."
    ),
    151: (
        "We processed all 272 published gross-removal tiles and assigned pixels to the four FAO GEZ biomes. Our area-weighted per-hectare means before excluding mangrove pixels are 0.960530, 1.376466, 1.430112 and 0.618633 tC ha⁻¹ yr⁻¹ for tropical, subtropical, temperate and boreal forest, respectively. These are our estimates from the Gibbs/GFW rasters, not rates directly reported by Gibbs et al."
    ),
    152: (
        "The matching forest-type tiles of the WRI/GFW model classify mangrove-factor pixels as code 6. We excluded those pixels from both the weighted gross-removal sum and the valid-pixel area before computing the rates used in the country account. The resulting non-mangrove rates are 0.934399, 1.370738, 1.430069 and 0.618633 tC ha⁻¹ yr⁻¹, respectively. The full pixel areas and the comparison with unmasked rates appear in Supplementary Section S3. The source-model classification does not establish the exact overlap between mangroves and FRA 2025 forest area."
    ),
    153: (
        "Each biome mean weights pixel rates by spherical pixel area, accounting for their latitude-dependent size in the geographic raster. We keep tropical and subtropical domains separate. The means are fixed in the main country calculation; pixel-to-pixel standard deviations are reported as spatial dispersion and used only in a sensitivity analysis, not interpreted as confidence intervals for the means."
    ),
    154: (
        "For country i, Cᵢ = Σᵦ Aᵢᵦ φᵦ and Uᵢ = (44/12) Cᵢ, where Aᵢᵦ is its allocated non-mangrove forest area in biome b (ha), φᵦ is the corresponding rate above (tC ha⁻¹ yr⁻¹), Cᵢ is carbon added to living biomass (tC yr⁻¹), and Uᵢ is the associated CO₂ quantity (tCO₂ yr⁻¹)."
    ),
    155: (
        "Because national uptake is estimated as area multiplied by an imposed biome-average rate, uptake per hectare is an estimated intensity, not an observed national efficiency or a measured net ecosystem sink."
    ),
    171: (
        "Uncertainty is propagated with 20,000 Monte Carlo draws of the GSCC. The four estimated Gibbs/GFW biome rates and FRA forest areas are fixed in the main calculation; the same GSCC draw is used for GCW and BCW so their values are combined draw by draw. Gibbs et al. (2025) report uncertainties for global and climate-domain gross-removal totals, not for our four GEZ per-hectare rates. We therefore do not transfer their standard deviations to the country estimates. Spatial pixel dispersion is shown separately as a deterministic sensitivity."
    ),
    173: (
        "These monetary intervals do not propagate uncertainty in FRA forest area, biome assignment, the Gibbs/GFW per-hectare rates or missing deforestation reports. They should be interpreted as conditional on those inputs."
    ),
    177: (
        "The Supplementary Material reports sensitivities to alternative biome shares, spatial dispersion around the Gibbs/GFW biome means, the mangrove-area convention, the fraction of full GSCC value assigned to forest additions, and the FAOSTAT net forest-stock-change comparator. FAOSTAT and Pan et al. (2024) describe different flux boundaries and are not inputs to the main Gibbs/GFW estimate."
    ),
    178: (
        "As a check on raster aggregation, the four unmasked GEZ pixel totals sum to 14.446 Gt CO₂ yr⁻¹, close to the rounded 14.5 Gt CO₂ yr⁻¹ gross-removal total published by Gibbs et al. (2025). Excluding source-model mangrove pixels gives 14.210 Gt CO₂ yr⁻¹ on that pixel support. The country estimate subsequently applies the non-mangrove rates to FRA 2025 national forest areas; it is not intended to reproduce the source-model global total exactly. The differing forest masks, biome definitions and reference years are documented as limitations."
    ),
    180: (
        "Three reference periods coexist. The Gibbs/GFW forest gross-removal rates average 2001-2023; FRA forest areas refer to 2025; population, GDP, emissions and debt refer to 2023 or the nearest available year compatible with the published blue carbon account. Reported deforestation is averaged over 2010-2019. This pairing supports comparative accounting but is not a contemporaneous measurement of every country's forest flow."
    ),
}

REFERENCE = (
    "Gibbs, D. A., Rose, M., Grassi, G., Melo, J., Rossi, S., Heinrich, V., & Harris, N. L. (2025). "
    "Revised and updated geospatial monitoring of 21st century forest carbon fluxes. "
    "Earth System Science Data, 17, 1217-1243. https://doi.org/10.5194/essd-17-1217-2025"
)


def main() -> None:
    if BACKUP.exists():
        if hashlib.sha256(BACKUP.read_bytes()).digest() != hashlib.sha256(PAPER.read_bytes()).digest():
            raise FileExistsError(f"Backup differs from current PAPER; refusing to overwrite: {BACKUP}")
    document = Document(PAPER)
    paragraphs = document.paragraphs
    if "Pan et al. (2024)" not in paragraphs[149].text or "Biome uptake rates" != paragraphs[148].text:
        raise RuntimeError("PAPER methods have changed; review paragraph mapping before editing")
    if "Gibbs, D. A., Rose, M." in "\n".join(p.text for p in paragraphs):
        raise RuntimeError("Gibbs citation already present; refusing a duplicate edit")
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists():
        shutil.copy2(PAPER, BACKUP)
    for index, value in REPLACEMENTS.items():
        paragraphs[index].text = value
    reference_paragraph = paragraphs[127].insert_paragraph_before(REFERENCE)
    reference_paragraph.style = paragraphs[126].style

    candidate = PAPER.with_name("PAPER_gibbs_methods_candidate.docx")
    document.save(candidate)
    with zipfile.ZipFile(BACKUP) as before, zipfile.ZipFile(candidate) as after:
        for name in ("word/comments.xml", "word/numbering.xml"):
            if name in before.namelist():
                original = etree.fromstring(before.read(name))
                revised = etree.fromstring(after.read(name))
                if etree.tostring(original, method="c14n") != etree.tostring(revised, method="c14n"):
                    raise AssertionError(f"Unexpected content change to {name}; candidate kept for review")
    checked = Document(candidate)
    text = "\n".join(p.text for p in checked.paragraphs)
    for token in ("14.5 ± 7.7", "0.934399", "0.960530", "Gibbs, D. A., Rose, M."):
        if token not in text:
            raise AssertionError(f"Missing expected method text: {token}")
    candidate.replace(PAPER)
    print(f"Updated {PAPER}")
    print(f"Backup {BACKUP}")


if __name__ == "__main__":
    main()
