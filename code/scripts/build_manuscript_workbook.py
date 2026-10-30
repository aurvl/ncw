"""Prepare the agreed manuscript workbook from pipeline tables and sources.

Notebook 09 is the user-facing entry point: it previews these frames and writes
``results/tables/manuscript/Manuscript_tables.xlsx``. This helper contains only
the reusable workbook assembly and formatting logic.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


BIOMES = ("tropical", "subtropical", "temperate", "boreal")
BCES = ("saltmarsh", "seagrass", "mangrove")
STATS = ("mean", "sd", "p05", "p50", "p95")

SHEET_DESCRIPTIONS = [
    ("00 Index", "Workbook index and reading order."),
    ("01 Table 1 - GCW by biome", "Forest area, Gibbs/GFW rate, annual additions and GCW by biome."),
    ("02 Table 2 - Marine contrib.", "Changes in positive removal-emissions balances after marine removals are added."),
    ("03 Table 3 - Forest loss", "Latest reported FRA forest loss and associated annual lost capacity by group."),
    ("04 Table 4 - 2% allocation", "Illustrative allocation equal to 2% of annual GCW by economic group."),
    ("05 Region profile", "Aggregated accounts and related indicators by region."),
    ("06 Economic group profile", "Aggregated accounts and related indicators by economic group."),
    ("08 Gibbs climate-domain rates", "Gibbs et al. published climate-domain and global removals totals."),
    ("09 Gibbs GEZ rates", "Four GEZ rates with and without source-model mangrove pixels."),
    ("10 Gibbs rate sensitivity", "Biome rate plus/minus one spatial standard deviation and entity-level results."),
    ("11 BCE rates and BCP", "Coastal BCE rates and the separate ocean biological pump flux."),
    ("12 GSCC sensitivity", "World account values under alternative fractions valued with the GSCC."),
    ("13 FRA loss periods", "Reported FRA loss rates for each period and the selected latest period."),
    ("14 Restoration", "Entity and soil-scenario restoration costs, targets and available supply."),
    ("15 Spatial results", "Spatial model fits, coefficients, diagnostics and comparisons."),
    ("16 Sources", "Source registry used by the analysis."),
    ("17 Mangrove harmonisation", "Entity-level mangrove-area evidence and selected source."),
    ("18 Unit crosswalk", "Source-to-entity mapping used to harmonize the blue-carbon records."),
]


def _csv(root: Path, relative: str) -> pd.DataFrame:
    path = root / relative
    if not path.is_file():
        raise FileNotFoundError(f"Required workbook input is missing: {relative}")
    return pd.read_csv(path, low_memory=False)


def _summary(values: np.ndarray) -> dict[str, float]:
    values = np.asarray(values, dtype=np.float64).ravel()
    values = values[np.isfinite(values)]
    if not values.size:
        return {stat: np.nan for stat in STATS}
    p05, p50, p95 = np.percentile(values, [5, 50, 95])
    return {"mean": float(np.mean(values)), "sd": float(np.std(values, ddof=1)) if values.size > 1 else 0.0,
            "p05": float(p05), "p50": float(p50), "p95": float(p95)}


def _sum_draws(archive: np.lib.npyio.NpzFile, key: str) -> np.ndarray:
    values = np.asarray(archive[key], dtype=np.float64)
    if values.ndim != 2:
        raise ValueError(f"Expected an entity-by-draw array for {key}; received {values.shape}.")
    matrix = values
    total = np.nansum(matrix, axis=0)
    total[np.all(~np.isfinite(matrix), axis=0)] = np.nan
    return total


def _group_table(frame: pd.DataFrame, group_col: str, columns: dict[str, str]) -> pd.DataFrame:
    available = [column for column in columns if column in frame]
    renamed = frame[[group_col, *available]].rename(columns={group_col: "economic_group", **columns})
    grouped = renamed.groupby("economic_group", dropna=False, sort=True)
    rows = []
    for label, group in grouped:
        row = {"economic_group": "Unclassified" if pd.isna(label) else label, "n_entities": len(group)}
        for col in group.columns.drop("economic_group"):
            values = pd.to_numeric(group[col], errors="coerce")
            row[col] = values.sum(min_count=1)
        rows.append(row)
    return pd.DataFrame(rows)


def build_workbook_tables(root: Path) -> dict[str, pd.DataFrame]:
    root = Path(root).resolve()
    table_root = root / "results" / "tables"
    manuscript_root = table_root / "manuscript"
    gcw = _csv(root, "results/tables/02_gcw_by_entity.csv")
    ncw = _csv(root, "results/tables/03_country_level_ncw.csv")
    loss = _csv(root, "results/tables/04_deforestation_by_entity.csv")
    restoration = _csv(root, "results/tables/05_restoration_costs_by_entity.csv")
    curves = _csv(root, "results/tables/06_restoration_supply_curves.csv")
    spatial = _csv(root, "results/tables/09_spatial_model_results.csv")

    # Table 1: global forest area, rate, gross additions, and value by biome.
    rates = _csv(root, "data/derived/gibbs_gfw/gfw_gibbs_biome_rates_non_mangrove.csv").set_index("biome")
    with np.load(root / "data/derived/draws/gcw_draws.npz", allow_pickle=True) as archive:
        gscc_draws = np.asarray(archive["gscc"], dtype=np.float64)
        global_gcw_draws = _sum_draws(archive, "gcw_usd2015_yr")
    table1_rows = []
    for biome in BIOMES:
        area = pd.to_numeric(gcw[f"forest_area_{biome}_ha"], errors="coerce").sum(min_count=1)
        tc = pd.to_numeric(gcw[f"forest_additions_{biome}_tc_yr"], errors="coerce").sum(min_count=1)
        tco2 = pd.to_numeric(gcw[f"forest_additions_{biome}_tco2_yr"], errors="coerce").sum(min_count=1)
        wealth = _summary(tco2 * gscc_draws)
        rate = rates.loc[biome]
        table1_rows.append({
            "biome": biome, "forest_area_ha": area,
            "forest_share_of_total": area / gcw["forest_area_total_ha"].sum(min_count=1),
            "gibbs_gfw_rate_mean_tc_ha_yr": rate["rate_tc_ha_yr_mean"],
            "gibbs_gfw_rate_spatial_sd_tc_ha_yr": rate["rate_tc_ha_yr_spatial_sd"],
            "forest_additions_tc_yr": tc, "forest_additions_tco2_yr": tco2,
            **{f"gcw_usd2015_yr_{stat}": wealth[stat] for stat in STATS},
        })
    gcw_total = _summary(global_gcw_draws)
    table1_rows.append({
        "biome": "Total GCW", "forest_area_ha": gcw["forest_area_total_ha"].sum(min_count=1),
        "forest_share_of_total": 1.0,
        "gibbs_gfw_rate_mean_tc_ha_yr": np.nan,
        "gibbs_gfw_rate_spatial_sd_tc_ha_yr": np.nan,
        "forest_additions_tc_yr": gcw["forest_additions_total_tc_yr"].sum(min_count=1),
        "forest_additions_tco2_yr": gcw["forest_additions_total_tco2_yr"].sum(min_count=1),
        **{f"gcw_usd2015_yr_{stat}": gcw_total[stat] for stat in STATS},
    })
    table1 = pd.DataFrame(table1_rows)

    # Table 2: positive balance counts on the matched emissions/account sample.
    balances = _csv(root, "results/tables/manuscript/Supplementary_accounting_balances_2024_by_unit.csv")
    group_col = "groups_x" if "groups_x" in balances else "groups_y"
    balances[group_col] = balances[group_col].fillna(balances.get("groups_y"))
    balance_rows = []
    for group, frame in balances.groupby(group_col, dropna=False, sort=True):
        green_positive = frame["gcw_accounting_balance_mean_tco2_yr"].gt(0)
        blue_positive = frame["ncw_accounting_balance_mean_tco2_yr"].gt(0)
        balance_rows.append({
            "economic_group": "Unclassified" if pd.isna(group) else group,
            "matched_entities": frame["UNIT_CODE"].nunique(),
            "positive_with_gcw": int(green_positive.sum()),
            "positive_with_ncw": int(blue_positive.sum()),
            "newly_positive_after_marine_additions": int((~green_positive & blue_positive).sum()),
            "net_change_in_balance_tco2_yr": pd.to_numeric(
                frame["ncw_accounting_balance_mean_tco2_yr"] - frame["gcw_accounting_balance_mean_tco2_yr"],
                errors="coerce",
            ).sum(min_count=1),
        })
    table2 = pd.DataFrame(balance_rows)

    # Table 3: FRA latest-period losses and corresponding forest-capacity values.
    table3_rows = []
    for group, frame in loss.groupby("economic_group", dropna=False, sort=True):
        reporters = frame["reported_deforestation"].fillna(False).astype(bool)
        positive = pd.to_numeric(frame["annual_deforestation_ha_yr"], errors="coerce").gt(0)
        table3_rows.append({
            "economic_group": "Unclassified" if pd.isna(group) else group,
            "entities": frame["unit_code"].nunique(),
            "reporting_entities": int(reporters.sum()),
            "entities_with_positive_reported_loss": int((reporters & positive).sum()),
            "annual_deforestation_ha_yr": pd.to_numeric(frame["annual_deforestation_ha_yr"], errors="coerce").sum(min_count=1),
            "lost_capacity_tco2_yr": pd.to_numeric(frame["lost_capacity_tco2_yr"], errors="coerce").sum(min_count=1),
            "lost_value_usd2015_yr_mean": pd.to_numeric(frame["lost_value_usd2015_yr_mean"], errors="coerce").sum(min_count=1),
            "median_erosion_pct_of_gcw": pd.to_numeric(frame["erosion_pct_of_gcw"], errors="coerce").median(),
            "selected_fra_periods": "; ".join(sorted(frame["deforestation_period"].dropna().astype(str).unique())),
        })
    table3 = pd.DataFrame(table3_rows)

    # Table 4: aggregate the stipulated 2% annual allocation by economic group.
    allocation_rows = []
    for group, frame in ncw.groupby("economic_group", dropna=False, sort=True):
        allocation = pd.to_numeric(frame["allocation_2pct_gcw_usd2015_yr"], errors="coerce")
        allocation_rows.append({
            "economic_group": "Unclassified" if pd.isna(group) else group,
            "entities": frame["unit_code"].nunique(),
            "entities_with_gcw": int(frame["has_gcw"].fillna(False).astype(bool).sum()),
            "annual_allocation_usd2015_yr": allocation.sum(min_count=1),
            "share_of_total_allocation_pct": 100 * allocation.sum(min_count=1) / pd.to_numeric(ncw["allocation_2pct_gcw_usd2015_yr"], errors="coerce").sum(min_count=1),
            "median_allocation_per_entity_usd2015_yr": allocation.median(),
        })
    table4 = pd.DataFrame(allocation_rows)

    profile_path = table_root / "07_geographic_economic_profiles.xlsx"
    profiles = pd.read_excel(profile_path, sheet_name=None)
    region_profile = profiles["Region"]
    group_profile = profiles["Economic group"]

    climate_domain = pd.DataFrame([
        {"climate_domain": "Tropical", "gross_removals_gtco2_yr": 7.1, "published_sd_gtco2_yr": 7.6, "period": "2001-2023", "sign_convention": "positive magnitude", "source": "Gibbs et al. (2025), Table 4"},
        {"climate_domain": "Subtropical", "gross_removals_gtco2_yr": 1.7, "published_sd_gtco2_yr": 0.56, "period": "2001-2023", "sign_convention": "positive magnitude", "source": "Gibbs et al. (2025), Table 4"},
        {"climate_domain": "Temperate", "gross_removals_gtco2_yr": 3.1, "published_sd_gtco2_yr": 0.55, "period": "2001-2023", "sign_convention": "positive magnitude", "source": "Gibbs et al. (2025), Table 4"},
        {"climate_domain": "Boreal", "gross_removals_gtco2_yr": 2.5, "published_sd_gtco2_yr": 0.95, "period": "2001-2023", "sign_convention": "positive magnitude", "source": "Gibbs et al. (2025), Table 4"},
        {"climate_domain": "Global", "gross_removals_gtco2_yr": 14.5, "published_sd_gtco2_yr": 7.7, "period": "2001-2023", "sign_convention": "positive magnitude", "source": "Gibbs et al. (2025), Table 4"},
    ])
    gez_parts = []
    for variant, filename in [
        ("including_source_model_mangrove_pixels", "gfw_gibbs_biome_rates_with_mangroves.csv"),
        ("excluding_source_model_mangrove_pixels", "gfw_gibbs_biome_rates_non_mangrove.csv"),
    ]:
        frame = _csv(root, f"data/derived/gibbs_gfw/{filename}").copy()
        frame.insert(0, "pixel_variant", variant)
        gez_parts.append(frame)
    gez_rates = pd.concat(gez_parts, ignore_index=True)

    sensitivity_rates = _csv(root, "data/derived/gibbs_gfw/gfw_gibbs_biome_rates_non_mangrove.csv")
    sensitivity_rows = []
    for _, row in sensitivity_rates.iterrows():
        mean = float(row["rate_tc_ha_yr_mean"])
        sd = float(row["rate_tc_ha_yr_spatial_sd"])
        for scenario, rate in [("central", mean), ("spatial_sd_low", max(0.0, mean - sd)), ("spatial_sd_high", mean + sd)]:
            sensitivity_rows.append({"scope": "biome_rate", "biome": row["biome"], "iso3": pd.NA,
                                    "scenario": scenario, "rate_tc_ha_yr": rate, "uptake_tco2_yr": np.nan})
    entity_sensitivity = _csv(root, "data/derived/forest_account/forest_uptake_spatial_sensitivity.csv")
    entity_sensitivity = entity_sensitivity.rename(columns={"iso3": "iso3", "scenario": "scenario", "uptake_tco2_yr": "uptake_tco2_yr"})
    entity_sensitivity.insert(0, "scope", "entity_total")
    entity_sensitivity["biome"] = pd.NA
    entity_sensitivity["rate_tc_ha_yr"] = np.nan
    rate_sensitivity = pd.concat([pd.DataFrame(sensitivity_rows), entity_sensitivity[
        ["scope", "biome", "iso3", "scenario", "rate_tc_ha_yr", "uptake_tco2_yr"]
    ]], ignore_index=True)

    with np.load(root / "data/raw/environmental/marine/bcw/output/draws.npz", allow_pickle=True) as archive:
        bce_draws = np.asarray(archive["bce_rates_draws"], dtype=np.float64)
        bcp = _csv(root, "data/raw/environmental/marine/bcw/output/country_level_bcw.csv")["BCP Seq (tC)"].dropna()
        bce_rows = []
        for index, name in enumerate(("saltmarsh", "seagrass", "mangrove")):
            values = _summary(bce_draws[index])
            bce_rows.append({"component": name, "measure": "sequestration_rate", "unit": "tC/km2/year", **values,
                             "source": "BCW ecosystem-rate draws"})
        bce_rows.append({"component": "ocean_biological_carbon_pump", "measure": "reported_flux_sum",
                         "unit": "tC/year", "mean": float(bcp.sum()), "sd": np.nan,
                         "p05": np.nan, "p50": np.nan, "p95": np.nan,
                         "source": "BCW source account; pump is a flux, not a per-area rate"})
        bce_rates = pd.DataFrame(bce_rows)

    gscc_sensitivity = _csv(root, "data/derived/validation/gcw_valuation_fraction_sensitivity_world.csv")
    loss_periods = loss[[
        "unit_code", "iso3", "country_or_territory", "economic_group",
        "loss_area_2010_2015_ha_yr", "loss_area_2015_2020_ha_yr", "loss_area_2020_2025_ha_yr",
        "deforestation_period", "annual_deforestation_ha_yr", "reported_deforestation",
    ]]
    source_registry = _csv(root, "data/metadata/source_registry.csv")
    mangrove_path = manuscript_root / "Supplementary_mangrove_reconciliation.csv"
    mangrove = pd.read_csv(mangrove_path, low_memory=False) if mangrove_path.is_file() else ncw[[
        "unit_code", "iso3", "country_or_territory", "mangrove_area_ha", "mangrove_area_source",
    ]]
    crosswalk = _csv(root, "data/derived/bcw_territory_mapping.csv")

    index = pd.DataFrame([{"order": i + 1, "sheet": name, "content": description}
                          for i, (name, description) in enumerate(SHEET_DESCRIPTIONS)])
    return {
        "00 Index": index,
        "01 Table 1 - GCW by biome": table1,
        "02 Table 2 - Marine contrib.": table2,
        "03 Table 3 - Forest loss": table3,
        "04 Table 4 - 2% allocation": table4,
        "05 Region profile": region_profile,
        "06 Economic group profile": group_profile,
        "08 Gibbs climate-domain rates": climate_domain,
        "09 Gibbs GEZ rates": gez_rates,
        "10 Gibbs rate sensitivity": rate_sensitivity,
        "11 BCE rates and BCP": bce_rates,
        "12 GSCC sensitivity": gscc_sensitivity,
        "13 FRA loss periods": loss_periods,
        "14 Restoration": restoration,
        "15 Spatial results": spatial,
        "16 Sources": source_registry,
        "17 Mangrove harmonisation": mangrove,
        "18 Unit crosswalk": crosswalk,
    }


def write_workbook(tables: dict[str, pd.DataFrame], output_path: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, frame in tables.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
            sheet = writer.book[sheet_name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            if len(frame.columns):
                table_ref = f"A1:{sheet.cell(row=max(len(frame) + 1, 2), column=len(frame.columns)).coordinate}"
                if len(frame):
                    table = Table(displayName=f"Table_{len(writer.book.sheetnames):02d}", ref=table_ref)
                    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                                          showColumnStripes=False)
                    sheet.add_table(table)
                for cell in sheet[1]:
                    cell.fill = PatternFill("solid", fgColor="18324B")
                    cell.font = Font(color="FFFFFF", bold=True)
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                sheet.row_dimensions[1].height = 30
                for column_cells in sheet.columns:
                    letter = column_cells[0].column_letter
                    width = min(48, max(12, max(len(str(cell.value or "")) for cell in column_cells[:250]) + 2))
                    sheet.column_dimensions[letter].width = width
    return output_path
