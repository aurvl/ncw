"""Create the BNC2 supplementary material from the current accounting protocol."""

from __future__ import annotations

from pathlib import Path
import json

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
import pandas as pd


OUTPUT = Path(__file__).resolve().parents[3] / "Supplementary Material.docx"
REPO = Path(__file__).resolve().parents[2]
BLACK = RGBColor(0, 0, 0)
GRAY = "D9D9D9"


def set_cell_shading(cell: object, fill: str) -> None:
    """Apply a solid cell background colour."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell: object, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    """Add readable internal padding to a table cell."""
    tc_mar = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        cell._tc.get_or_add_tcPr().append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    """Add a compact, readable table with a dark header row."""
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = label
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "404040")
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[index])
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(8)
    document.add_paragraph()


def add_paragraph(
    document: Document,
    text: str,
    style: str | None = None,
    keep_together: bool = False,
) -> None:
    """Add one body paragraph with the project-wide spacing rules."""
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    paragraph.paragraph_format.keep_together = keep_together
    paragraph.add_run(text)


def remove_paragraph_borders(paragraph: object) -> None:
    """Remove inherited Word paragraph borders from a title paragraph."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def configure_document(document: Document) -> None:
    """Apply a restrained research-paper layout."""
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    styles = document.styles
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = BLACK
    title_ppr = styles["Title"]._element.pPr
    if title_ppr is not None:
        title_borders = title_ppr.find(qn("w:pBdr"))
        if title_borders is not None:
            title_ppr.remove(title_borders)
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.size = Pt(16)
    styles["Heading 1"].font.size = Pt(13)
    styles["Heading 2"].font.size = Pt(11)
    styles["Heading 3"].font.size = Pt(10)


def build_document() -> Document:
    """Build the new supplementary material in the order used by the paper."""
    rates = pd.read_csv(REPO / "data" / "outputs" / "gibbs_gfw" / "gfw_gibbs_biome_rates.csv")
    rates_with_mangroves = pd.read_csv(
        REPO / "data" / "outputs" / "gibbs_gfw" / "gfw_gibbs_biome_rates_with_mangroves.csv"
    )[["biome", "rate_tc_ha_yr_mean"]].rename(
        columns={"rate_tc_ha_yr_mean": "rate_with_mangroves"}
    )
    rates = rates.merge(rates_with_mangroves, on="biome", validate="one_to_one")
    valuation_fractions = pd.read_csv(
        REPO / "data" / "outputs" / "validation" / "gcw_valuation_fraction_sensitivity_world.csv"
    )
    finance = json.loads(
        (REPO / "data" / "outputs" / "validation" / "finance_action_cost_comparison.json").read_text(encoding="utf-8")
    )
    busch_status = pd.read_csv(
        REPO / "data" / "outputs" / "validation" / "busch_action_curve_status_summary.csv"
    )
    busch_status_labels = {
        "no_country_curve_in_published_table": "No national curve available",
        "target_exceeds_published_supply_at_100": "Target exceeds supply at $100/tCO2",
        "target_within_positive_price_curve": "Target within positive-price range",
        "target_within_supply_at_nonpositive_price": "Target within zero/negative-price range",
        "target_within_supply_at_or_below_minus20": "At/below −$20/tCO2 (subset)",
    }
    allocation = pd.read_csv(
        REPO / "data" / "outputs" / "validation" / "our_2pct_gcw_allocation_and_pledge.csv"
    )
    allocation_column = "our_2pct_gcw_allocation_usd2015_yr"
    allocation_ranked = allocation.sort_values("gcw_usd2015_yr_mean", ascending=False)
    allocation_total = float(allocation[allocation_column].sum())
    top10_share = float(allocation_ranked.head(10)[allocation_column].sum() / allocation_total * 100)
    top50_share = float(allocation_ranked.head(50)[allocation_column].sum() / allocation_total * 100)
    cell_reconstruction_dir = REPO / "data" / "outputs" / "busch_reconstruction"
    cell_reconstruction = json.loads(
        (cell_reconstruction_dir / "busch_mangrove_refresh_20261023_summary.json").read_text(encoding="utf-8")
    )
    cell_reconstruction_coverage = pd.read_csv(
        cell_reconstruction_dir / "busch_mangrove_refresh_20261023_coverage_105.csv"
    )
    cell_reconstruction_loss = float(cell_reconstruction_coverage["lost_capacity_tco2_yr"].sum())
    cell_reconstruction_within = cell_reconstruction_coverage.loc[
        cell_reconstruction_coverage["status"].eq("target_within_cell_supply"), "lost_capacity_tco2_yr"
    ]
    cell_reconstruction_loss_share = float(cell_reconstruction_within.sum() / cell_reconstruction_loss * 100)
    spatial_moran = pd.read_csv(
        REPO / "data" / "outputs" / "validation" / "spatial_emissions_2024_step1_moran.csv"
    )
    spatial_diagnostics = json.loads(
        (REPO / "data" / "outputs" / "validation" / "spatial_emissions_2024_diagnostics.json").read_text(encoding="utf-8")
    )
    gmw_comparison = pd.read_csv(REPO / "data" / "outputs" / "validation" / "gmw_2025_spatial_bcw_world.csv")
    deforestation_trend = pd.read_csv(
        REPO / "results" / "tables" / "manuscript" / "Supplementary_deforestation_trend.csv"
    )
    biome_order = ["tropical", "subtropical", "temperate", "boreal"]
    rates["biome"] = pd.Categorical(rates["biome"], categories=biome_order, ordered=True)
    rates = rates.sort_values("biome")
    document = Document()
    configure_document(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Forest and Marine Carbon Wealth Supplementary Material")
    remove_paragraph_borders(title)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Integrating forest and marine carbon accounts into the natural carbon wealth of nations")

    document.add_heading("S1 Accounting object", level=1)
    add_paragraph(
        document,
        "Green Carbon Wealth (GCW) is the annual monetary valuation of the forest carbon flow measured here: gross carbon additions to above- and below-ground living biomass. "
        "It is not a forest carbon stock, a market price, a finance requirement, or an estimate of additional mitigation. "
        "For each accounting unit, the annual flow in tonnes of carbon dioxide is valued with the Global Social Cost of Carbon (GSCC). "
        "The calculation does not estimate how long each addition remains stored and is not a net ecosystem carbon balance. "
        "Natural Carbon Wealth combines the forest and published marine components under the stated mangrove-area convention.",
    )
    add_paragraph(
        document,
        "All values use constant 2015 US dollars. Forest uptake rates are means for 2001-2023, applied to FRA 2025 forest areas. "
        "The deforestation analysis uses annual conversion reported for 2010-2020. "
        "Population and GDP use 2024 World Development Indicators; fossil and industrial CO2 emissions use the Global Carbon Budget 2025 country estimates for 2024. Debt is not an input to the current spatial model sequence. "
        "The emissions series excludes land-use change. This keeps it separate from the forest-conversion capacity-loss analysis.",
    )

    document.add_heading("S2 Data, accounting units, and harmonisation", level=1)
    add_table(
        document,
        ["Input", "Source and period", "Use in the analysis"],
        [
            ["Forest area", "FAO FRA 2025", "Country forest area; mangrove subtraction only where overlap with Forest is documented"],
            ["Mangroves", "GMW v4.1.12 map for 2025", "Mapped area used to update the marine account; spatially allocated with the study's marine boundaries"],
            ["Biome shares", "FAO GEZ 2010, country boundaries, JRC Global Forest Cover 2020", "Forest area by tropical, subtropical, temperate, and boreal biome"],
            ["Forest uptake rates", "Gibbs et al. 2025 / GFW, 2001-2023", "Gross additions to above- and below-ground living biomass, summarized by FAO GEZ biome"],
            ["Blue Carbon Wealth", "Hilmi et al. account, mangrove area updated with GMW 2025", "Coastal blue-carbon and oceanic biological-pump components; published rates retained"],
            ["Deforestation", "FAO FRA 2025", "Reported annual forest conversion over 2010-2020"],
            ["Socioeconomic indicators", "WDI 2024; Global Carbon Budget 2025, emissions year 2024", "GDP, population, and fossil and industrial CO2 emissions; debt is not used in the current spatial models"],
        ],
    )
    add_paragraph(
        document,
        "The accounting units follow FRA countries and territories. BCW rows are mapped to those units with a documented crosswalk. "
        "BCW sub-territories are regrouped where they belong to one FRA unit. Population and GDP are never summed across parts. "
        "Taiwan is included in China, consistent with FRA reporting. BCW-only territories and joint-regime rows remain separate where no FRA country match exists.",
    )

    forest_heading = document.add_heading("S3 Forest uptake estimate", level=1)
    forest_heading.paragraph_format.page_break_before = True
    add_paragraph(
        document,
        "Gibbs et al. (2025, Table 4) publish the following average annual gross-removal totals for 2001-2023 in their own climate domains. "
        "They use a negative sign for removals under an emissions-positive convention; magnitudes are shown here as positive quantities. "
        "These are total flows in Gt CO2/yr, not rates in tC/ha/yr, and their climate-domain boundaries are not identical to the FAO GEZ allocation used below.",
    )
    add_table(
        document,
        ["Gibbs climate domain", "Published gross removals (Gt CO2/yr)", "Published SD (Gt CO2/yr)"],
        [
            ["Tropical", "7.1", "7.6"],
            ["Subtropical", "1.7", "0.56"],
            ["Temperate", "3.1", "0.55"],
            ["Boreal", "2.5", "0.95"],
            ["Global", "14.5", "7.7"],
        ],
    )
    add_paragraph(
        document,
        "Our country estimate uses Route B. FRA 2025 forest area is distributed across tropical, subtropical, temperate, and boreal biomes using FAO GEZ shares. "
        "For each pixel p, the published cumulative gross-removal value R_p (Mg CO2/ha over 2001-2023) is converted to an annual carbon rate as r_p = R_p × (12/44) / 23. "
        "Within each GEZ biome, we weight r_p by the pixel's spherical area. The source-model forest-type raster identifies mangroves with code 6; those pixels are excluded from both the numerator and the area denominator of the rate applied to countries. "
        "The rates describe mean annual gross additions to above- and below-ground living forest biomass. They do not measure a net ecosystem carbon balance or carbon stored in soils, dead organic matter, or harvested wood, and they do not establish a residence time for the added carbon. "
        "The forest-area calculation separately subtracts FRA mangrove hectares documented as Forest: exact 100% Forest matches and the quantified 80% Forest share reported for Mozambique. Unresolved categories are not subtracted. The rule, its classification-stability assumption for Mozambique, and its limits are described in S4.",
    )
    add_table(
        document,
        ["FAO GEZ biome", "Our rate incl. mangroves (tC/ha/yr)", "Our rate excl. mangroves (tC/ha/yr)", "Valid non-mangrove pixel area (ha)"],
        [[
            str(row.biome),
            f"{row.rate_with_mangroves:.3f}",
            f"{row.rate_tc_ha_yr_mean:.3f}",
            f"{row.area_ha:,.0f}",
        ] for row in rates.itertuples(index=False)],
    )
    add_paragraph(
        document,
        "Both per-hectare columns are our estimates from Gibbs/GFW pixels; Gibbs et al. did not publish these four FAO GEZ means. Only the non-mangrove column enters the country calculation. "
        "The spatial standard deviation describes pixel-to-pixel dispersion within a biome; it is not an uncertainty interval for the biome mean. "
        "The published Gibbs standard deviations apply to the global and climate-domain totals above and cannot be transferred to our four rates. The country estimate applies each fixed non-mangrove mean to FRA forest area allocated to that biome.",
    )

    document.add_heading("S4 Mangroves and the forest-marine overlap", level=1)
    add_paragraph(
        document,
        "FRA forest area can include mangroves that meet the forest definition, whereas the separately reported FRA mangrove category can include forest and other wooded land. The global FRA table does not give a Forest/Other wooded land split for every unit. "
        "For GCW, mangrove hectares are subtracted when a national report documents a Forest share: this includes 2,082,920 ha in 12 exact 100% Forest units and 1,809,830 ha in 11 stability-assumption units, plus 287,848 ha in Mozambique and 312 ha in Mayotte from their reported Forest shares. The documented Forest overlap is 4,180,910 ha. The Mozambique, Mayotte and copied/forecast 2025 applications state their classification-stability assumptions. The evidence table also documents 72,477.75 ha as Other wooded land and 42.25 ha as Remaining land; neither is subtracted. PAK, SHN, TCA and the other units without a comparable quantitative split are not subtracted; their unknown Forest/Other wooded land area totals 11,613,440 ha. Kiribati's match relies on a desk-study source and has weaker evidential support. "
        "The current GMW extraction assigns 0 ha to Montserrat (MSR) and Ascension, Saint Helena and Tristan da Cunha (SHN), while FRA 2025 reports 10 ha and 2,000 ha, respectively, in its mangrove category. This cross-source coverage mismatch is not interpreted as confirmed absence of mangroves. "
        "For BCW, mapped mangrove area is updated using GMW v4.1.12 for 2025, the latest complete global annual extent series identified, released on 3 August 2026. GMW's global accuracy assessment reports an F1 score of 0.9276 (95% CI 0.9220-0.9330), with local accuracy varying by place. The update directly intersects the 2025 GMW polygons with the study's EEZ UNION geometries after reprojection to Mollweide, retains the source cntry_iso3 for provenance, and assigns area according to the intersected EEZ geometry. Following the GMW README, the raw geometric intersections are multiplied by its global 0.9775 area factor before aggregation. The 234 country/territory EEZ units are mapped to study units; overlapping EEZ geometries may cause physical-area overlap in the summed area. The published mangrove-burial rate and other marine components are retained. These area rules do not establish exact non-overlap of carbon flows for every country.",
        keep_together=True,
    )

    document.add_heading("S5 Uncertainty and aggregation", level=1)
    add_paragraph(
        document,
        "The analysis uses 20,000 Monte Carlo draws for the GSCC. Biome mean rates and forest areas are held fixed in the main estimate; spatial dispersion is reported separately as a deterministic sensitivity, not sampled as rate uncertainty. "
        "GCW and BCW use aligned GSCC draws. Global Table 1 reports available-component totals: GCW across its 236 units and BCW across its 179 country or territory accounts plus the source Joint Regimes residual. Global NCW is their sum on each aligned draw, giving US$10.590 trillion 2015/year from mean GCW of US$8.261 trillion and mean BCW of US$2.329 trillion. "
        "Unit-level NCW sums the GCW and BCW components available for each unit, treating an unavailable component as zero; units with neither component have no NCW value. Complete-account comparisons are reported separately for units with both accounts. "
        "Reported intervals are the 5th and 95th percentiles of the corresponding draw distribution. Standard deviations are not added across components.",
    )
    add_paragraph(
        document,
        "The resulting monetary intervals therefore reflect GSCC uncertainty and uncertainty in the blue account, not statistical uncertainty in the Gibbs biome rates, national forest area, or biome assignment.",
    )
    valuation_heading = document.add_heading("Conditional GCW valuation fraction sensitivity", level=2)
    valuation_heading.paragraph_format.page_break_before = True
    add_paragraph(
        document,
        "The table multiplies each GCW draw by an illustrative valuation fraction while retaining its GSCC draw. "
        "The fractions are assumptions about the share of the full-GSCC valuation applied to the measured annual additions to living biomass. "
        "They are not observed carbon-retention fractions and do not establish a storage duration or equivalence to a permanent removal.",
    )
    add_table(
        document,
        ["Valuation fraction", "Mean (billion 2015 US$/yr)", "P05", "P95"],
        [
            [
                f"{row.valuation_fraction:.0%}",
                f"{row.mean_billion_usd2015_yr:,.1f}",
                f"{row.p05_billion_usd2015_yr:,.1f}",
                f"{row.p95_billion_usd2015_yr:,.1f}",
            ]
            for row in valuation_fractions.itertuples(index=False)
        ],
    )

    document.add_heading("S6 Statistical analyses and supplementary results", level=1)
    add_paragraph(
        document,
        "R1 reports absolute GCW, GCW relative to GDP, concentration, and the correlation between GCW and GDP. "
        "The group comparison uses a Kruskal-Wallis test. Pairwise comparisons use Mann-Whitney tests with Holm adjustment. "
        "Per-capita indicators are reported in the supplementary country table because they do not distinguish the main green-carbon narrative across groups.",
    )
    add_paragraph(
        document,
        "R2 classifies units as green-dominant when GCW represents at least 66.7% of NCW and blue-dominant when it represents at most 33.3%. "
        "It compares forest-only and forest-plus-marine flow balances after subtracting 2024 fossil and industrial emissions in the 161-unit common sample. The forest term is gross addition to living biomass, so this comparison is not a net ecosystem carbon balance. Debt is not included in the current spatial model sequence.",
    )
    add_paragraph(
        document,
        "Country-level GCW and NCW results, deforestation outcomes, and the full set of derived descriptive and spatial statistics are retained in the project result files. "
        "The retention-valuation sensitivity is reported in S5 and in its accompanying CSV files.",
    )

    document.add_heading("S7 Annual forest-addition capacity no longer supplied after deforestation", level=1)
    add_paragraph(
        document,
        "R3 defines a loss as forest area converted each year multiplied by the country's allocated biome shares and the non-mangrove Gibbs/GFW gross-addition rates for the four FAO GEZ biomes. "
        "The result is the annual gross-addition capacity associated with converted area, no longer supplied while that area remains outside forest. It is not an estimate of direct land-use emissions or of carbon released by deforestation. "
        "The calculation retains missing reports as missing values. It does not treat them as zero loss.",
    )
    add_paragraph(
        document,
        "FRA reports deforestation for 139 of 236 forest accounting units. These units represent 78.7% of global GCW. "
        "The global loss totals therefore describe the reporting units only. The allocation of reported deforestation across national biomes is an assumption. "
        "It may understate tropical losses where conversion is concentrated in tropical forest within a multi-biome country.",
    )
    add_paragraph(
        document,
        "The accumulated shortfall scenario assumes converted area remains outside forest and applies the current mean biome rates; it does not model regrowth or changing rates. "
        "The trend analysis uses the five FRA reporting intervals. It reports Theil-Sen slopes and Mann-Kendall tests. "
        "Five time points provide limited power. Country trend results should therefore be read as descriptive evidence.",
    )
    add_paragraph(document, "Supplementary Table S3. Reported global deforestation by FRA reporting interval.")
    add_table(
        document,
        ["FRA reporting interval", "Interval midpoint", "Reported deforestation (Mha/yr)"],
        [[
            str(row.fra_interval),
            f"{row.midpoint_year:g}",
            f"{row.reported_deforestation_mha_per_year:.2f}",
        ] for row in deforestation_trend.itertuples(index=False)],
    )

    document.add_heading("S8 Climate finance scenarios", level=1)
    add_paragraph(
        document,
        "The main finance scenario applies our stipulated allocation rule to all 236 units with a GCW estimate: each unit is allocated 2% of its annual GCW value. "
        f"With the recalculated world GCW, this allocates {allocation_total / 1e9:.3f} billion constant 2015 US dollars per year. "
        f"The ten largest GCW units account for {top10_share:.2f}% of the total and the fifty largest for {top50_share:.2f}%. "
        "This is an allocation scenario, not an estimate of national need, observed climate-finance flows, or action costs.",
    )
    add_paragraph(
        document,
        "The separate illustration of a 100 billion US dollar annual pledge for 30 years totals 3 trillion dollars cumulatively. It is not directly subtracted from annual GCW, because the two quantities use different time units. "
        "The 2% rule remains an explicit assumption; it is not derived from the cost curves below.",
    )
    add_paragraph(
        document,
        "Busch et al. (2024) provide a complementary action-cost comparison. Their curves estimate the cost and supply of natural regeneration and plantation actions on eligible land, not the cost of protecting existing forest. We compare the curves with the deforestation-related capacity loss only as a stated scenario: the annual loss is assumed to persist for 30 years and is discounted at 5% to match the curve horizon.",
    )
    add_table(
        document,
        ["Busch curve status among positive-loss reporters", "Units", "Capacity loss (MtCO2/yr)"],
        [[
            busch_status_labels.get(str(row.status), str(row.status)),
            f"{int(row.countries)}",
            f"{row.lost_capacity_tco2_yr / 1e6:.3f}",
        ] for row in busch_status.itertuples(index=False)],
    )
    add_paragraph(
        document,
        f"Of {int(finance['positive_loss_reporters'])} units with reported positive loss, {int(finance['country_curves_available'])} have a country curve. "
        f"For {int(finance['matched_countries_supply_within_price_domain'])} units the discounted target falls within the published supply domain, representing {finance['all_reported_loss_supply_within_price_domain_share_pct']:.1f}% of reported capacity loss. "
        f"Only {int(finance['positive_price_curve_countries'])} units have a positive marginal-price bracket; its median bracket is {finance['positive_price_median_bracket_usd2020_tco2'][0]:.1f}-{finance['positive_price_median_bracket_usd2020_tco2'][1]:.1f} US$2020/tCO2. "
        f"The corresponding uniform-price proxy is {finance['positive_price_payment_proxy_usd2015_yr_partial_lo'] / 1e6:.2f}-{finance['positive_price_payment_proxy_usd2015_yr_partial_hi'] / 1e6:.2f} million US$2015/year for those units only. It is not an integrated project cost or a global financing need. "
        "The full unit-by-unit allocation and curve statuses are in the accompanying CSV outputs and the comprehensive results inventory.",
    )
    add_paragraph(
        document,
        f"A separate reconstruction of the public Busch cell outputs uses both workbook sheets and re-estimates country supply after converting source quantities from tC to tCO2 and counting the natural-regeneration soil term once. "
        f"It reproduces 71 source-country curves and places {int(cell_reconstruction['finance_target_within_supply'])} of the 105 positive-loss targets within reconstructed supply, covering {cell_reconstruction_loss_share:.1f}% of the sample's annual capacity loss ({cell_reconstruction_loss / 1e6:.3f} MtCO2/year in total). "
        f"Twelve targets exceed reconstructed supply and 34 have no matched source curve. These are our derived results, not values published by Busch et al.; the unit and soil corrections remain provisional because the reconstructed quantities require clarification against the published aggregate. "
        "They are reported separately from the published-grid screening above and are not used to revise its Figure 4 counts.",
    )

    document.add_heading("S9 Spatial patterns and exploratory models", level=1)
    spatial_moran = spatial_moran.loc[spatial_moran["measure"].eq("y_log1p")].copy()
    spatial_relations = [
        ("Capacity loss / base", "baseline_2024_same_sample", "lost_capacity_mtco2_yr"),
        ("Relative erosion / base", "baseline_2024_same_sample", "erosion_pct_of_gcw"),
        ("GCW/GDP / base", "baseline_2024_same_sample", "gcw_gdp_pct_2024"),
        ("Capacity loss / + emissions X", "add_emissions_2024_as_x", "lost_capacity_mtco2_yr"),
        ("Relative erosion / + emissions X", "add_emissions_2024_as_x", "erosion_pct_of_gcw"),
        ("GCW/GDP / + emissions X", "add_emissions_2024_as_x", "gcw_gdp_pct_2024"),
        ("Emissions Y / + capacity loss X", "emissions_2024_as_y_plus_lost_capacity_x", "emissions_mtco2_yr_2024"),
        ("Emissions Y / + erosion X", "emissions_2024_as_y_plus_erosion_x", "emissions_mtco2_yr_2024"),
        ("Emissions Y / + GCW/GDP X", "emissions_2024_as_y_plus_gcw_gdp_x", "emissions_mtco2_yr_2024"),
    ]
    rows = []
    for label, specification, outcome in spatial_relations:
        current = spatial_moran.loc[
            spatial_moran["specification"].eq(specification) & spatial_moran["outcome"].eq(outcome)
        ]
        row = {r.weight: r for r in current.itertuples(index=False)}
        q, k = row["queen_contiguity"], row["symmetric_knn5"]
        rows.append([
            label, str(int(q.n)),
            f"{q.moran_i:.3f} ({q.p_holm_within_specification_measure:.4f})",
            f"{k.moran_i:.3f} ({k.p_holm_within_specification_measure:.4f})",
        ])
    if len(rows) != 9:
        raise ValueError(f"Expected nine primary spatial relationships in the current Moran output; found {len(rows)}")
    add_paragraph(
        document,
        "The first three rows use each existing outcome as Y with base controls; the next three add 2024 emissions as X; the last three use emissions as Y and add one existing outcome as X. The nine relationships are tested under Queen and symmetric kNN-5 weights. The table reports Moran's I for the log-transformed outcome, with Holm-adjusted permutation p-values in parentheses. OLS residuals are also tested using 9,999 permutations.",
    )
    add_table(
        document,
        ["Relationship", "N", "Queen I (Holm p)", "kNN I (Holm p)"],
        rows,
    )
    add_paragraph(
        document,
        f"The primary set contains {spatial_diagnostics['number_of_primary_relationship_profiles']} relationships and {spatial_diagnostics['number_of_primary_relationship_weight_cases']} weight cases. With references and auxiliary specifications, the complete run covers {spatial_diagnostics['number_of_specification_outcome_weight_cases']} cases and {spatial_diagnostics['model_fit_rows']} fits (OLS, SAR, SEM, SDM, SDEM, and GNS), with {spatial_diagnostics['moran_rows']} Moran, {spatial_diagnostics['lm_rows']} LM, {spatial_diagnostics['nested_lr_rows']} nested likelihood-ratio, and {spatial_diagnostics['predictor_block_nested_tests']} predictor-block tests. "
        "Coefficient, fit-rank, convergence, sample, and diagnostic outputs are retained for review. No model is selected automatically; estimates are descriptive associations, not causal estimates.",
    )

    document.add_heading("S10 Limitations", level=1)
    add_paragraph(
        document,
        "The four mean Gibbs/GFW rates do not capture national differences in forest age, disturbance, management, or climate response. "
        "They measure gross additions to living biomass, not the full ecosystem carbon balance or storage duration.",
    )
    add_paragraph(
        document,
        "Reference periods differ: forest rates cover 2001-2023 and are applied to FRA 2025 area; socioeconomic inputs are from 2024; deforestation is averaged over 2010-2020. "
        "This temporal mismatch limits causal interpretation, and deforestation reporting is incomplete.",
    )
    add_paragraph(
        document,
        "GCW and NCW are annual welfare contributions, not carbon stocks, ownership, payment obligations, or additional mitigation. "
        "The finance scenarios are imposed allocations or action-cost comparisons; they do not measure observed transfers, funding needs, or the cost of protecting all forests.",
    )

    document.add_heading("S11 Data and code availability", level=1)
    add_paragraph(
        document,
        "Country estimates and validation tables were regenerated with `code/scripts/main.py`; extended tables use `code/scripts/build_manuscript_artifacts.py`. Notebooks 03-06 have been rerun against the current estimates. "
        "The four figures embedded in PAPER.docx have not yet been replaced and are handled as a separate manuscript synchronization step.",
        keep_together=True,
    )
    return document


def main() -> None:
    """Write the supplementary material to the project root."""
    document = build_document()
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
